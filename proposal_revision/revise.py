from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import shutil, json

root=Path(__file__).parent
src=Path('C:/Users/linke/Downloads/面向时空自适应稀疏建模的轻量化行人轨迹预测方法研究_开题报告草案.docx')
shutil.copy2(src,root/src.name)
d=Document(src)
ps=list(d.paragraphs)
changes={
13:'现有研究以平均位移误差（ADE）和最终位移误差（FDE）为主要精度指标，并通过社会交互、多模态生成与场景融合增强预测能力。面向轻量化部署，还需考察输入规模与实际执行开销：部分模型在完整的行人对关系矩阵上完成编码后再施加掩码，历史编码则通常使用固定间隔的完整序列。空间邻居与历史状态的贡献并不均衡，因此有必要研究信息选择能否减少后续计算，并通过实测检验其精度与效率代价。',
15:'从时间维度看，长期历史可提供运动趋势，近期状态可描述转向、加减速等局部变化，但各类信息的相对价值依赖场景与运动模式。前期高频历史实验显示，提高输入时间分辨率能够改善当前模型的预测表现；延长历史窗口的总体收益则较小。这为时间状态选择提供了研究依据，但不能据此直接认定远期状态无用，或认定近期密集采样优于均匀采样。需要在相同历史覆盖和固定输入预算下，检验哪些状态应被优先保留。',
27:'轻量化预测需要同时关注参数规模、输入信息量和计算执行方式。限制邻域、压缩输入状态与采用并行解码，分别作用于不同计算环节，其收益取决于模型结构、数据分布和运行硬件。结合上述研究，本课题重点考察历史轨迹内部的状态选择：在固定输入预算下，如何兼顾近期动态、运动事件和长期覆盖，并将其与空间交互稀疏化共同评价。',
28:'结构稀疏并不自动带来实际加速。在完整关系张量上计算后施加掩码，通常不能消除此前的关系编码开销；按边执行虽然可减少算术量，却可能引入索引、数据搬运和不规则归约成本。本课题据此将图结构、实际计算量、显存占用与端到端延迟分别验证，研究固定宽度邻居布局能否改善执行规则性，并量化填充槽位带来的额外计算。',
37:'综上，社会交互、时间依赖与多模态生成已形成较丰富的研究基础。本课题聚焦其中两个可检验问题：其一，自适应邻居选择能否在保留预测能力的同时减少关系计算，并在何种输入规模下体现执行收益；其二，在固定历史覆盖和时间状态预算下，非均匀采样能否比均匀采样更有效地保留运动信息。研究以空间交互稀疏化与时间状态选择为两个创新方向，以现有并行非自回归解码器为统一后端，通过受控实验建立精度与实际资源开销之间的对应关系。',
39:'本课题已建立轨迹预处理、训练、候选评估与效率分析流程，形成以时序—社会交叉编码器为核心的当前基线（项目编号 F16）。模型使用历史二维轨迹及派生运动状态，不依赖图像或高精地图分支；以相对位置、相对速度、接近趋势（closing）、预计碰撞时间（TTC）和最近接近距离（DCA）等关系量描述潜在交互。正式配置采用运动学自适应有向图，非自环邻居预算范围为 2～6，实际有效邻居数受可用行人与掩码约束。目标行人通过交叉注意力访问选中邻居的历史状态，后端以 20 个模式查询和终点／路径形状锚点并行生成候选轨迹。',
40:'空间方向已完成按边执行（EdgeExec）的数值等价与资源开销审计（EXP-197）。在同一模型检查点下，20 个批次的图结构不一致数为 0，最终轨迹最大绝对误差为 1.91×10⁻⁶，满足既定数值容差；模型总参数量保持为 355,048。在 ETH/UCY 输入的相同运行环境下，实际行人对数量相对完整矩阵减少约 58%～63%，分析器可统计的完整模型浮点运算量减少约 9%～10%。在固定 4 个非自环邻居并保留自环的合成扩展测试中，N=64 时峰值已分配显存由 30.42 MB 降至 4.15 MB，N=128 时由 152.45 MB 降至 9.21 MB（1 MB=10⁶ 字节）。这些结果支持关系计算与显存扩展方面的收益，不能直接等同于端到端加速。',
41:'EXP-197 的真实输入计时显示，批量大小为 1 时两种执行路径的平均延迟基本持平，为 64 时 EdgeExec 平均延迟增加约 6.85%；在已测试的合成规模中也未出现稳定的延迟优势。不规则索引与归约是待验证的开销来源。后续固定宽度执行（Fixed-K，EXP-198）已完成实现及 20 个批次的等价检查，轨迹最大绝对误差为 2.38×10⁻⁶。为保持正式图语义，执行总宽度应为 7，即最多 6 个非自环邻居加 1 个自环，而非直接设为 4。当前槽位利用率约 42.7%；已有计时记录波动较大，实际加速结论仍需在稳定运行条件下复测。',
42:'时间方向已由初始桥接实验（EXP-195）推进至训练预算受控与三随机种子复验（EXP-196）。AV2 行人诊断子集使用 8,000 个训练目标、2,000 个评估目标，统一预测未来 12 步、2.5 Hz（4.8 s）；三组分别输入 8 步／2.5 Hz、16 步／5 Hz 和 32 步／10 Hz 的历史。三组从同一检查点续训 4 个轮次，消除了初始实验中的额外续训混杂。三随机种子的 minADE20 均值±标准差依次为 0.2041±0.0008、0.2001±0.0002、0.1960±0.0003 m，selected ADE 依次为 0.4772、0.4708、0.4537 m。10 Hz 相对 2.5 Hz 的 minADE20 改善在各随机种子的场景级配对重采样中均获得不跨零的 95% 置信区间。该结果支持当前协议下高频输入的价值，但输入 Token 数同时增加，尚不能证明固定预算时间选择有效。',
43:'进一步的单随机种子诊断中，历史由名义 3.2 s 延长至 5.0 s，minADE20 改善约 1.38%，改善率的 95% 置信区间为 −0.45%～3.03%；转向组出现更明显的描述性改善，仍需复验。注意力在近期历史上较集中仅提供选点线索，不能作为删除远期状态的因果证据。此外，当前 2,000 个 AV2 评估目标中，6 m 内无其他行人的比例为 84.55%，密集交互样本有限，因此空间稀疏性的验证还需结合 ETH/UCY、SDD 与合成规模测试，避免将该子集的结论外推至拥挤场景。',
73:'执行层面比较掩码稠密执行、按边执行与固定宽度执行三条路径。令 N 为行人数、T 为历史状态数、E 为选中边数、W 为含自环的物理槽位宽度，在特征维度固定时，关系编码与时序—社会交叉注意力的主要行人对—时间计算由 O(N²T) 转为按边的 O(ET)，或规则布局下的 O(NWT)。正式配置取 W=7，并对不足的槽位施加有效性掩码。该复杂度结论仅针对相应交互分支；构图可能仍含 O(N²) 成本，时间自注意力、解码与数据搬运也须计入整体开销。',
75:'本研究将进入主编码器的单个历史状态称为时间 Token，并以 M 表示固定预算。首先在同一高频历史缓冲区、相同起止覆盖范围下比较三种等预算策略：均匀采样（Uniform）、近期密集采样（Recent-Dense）与远期密集采样（Distant-Dense）。三者保留相同数量的不同时间点，均包含当前时刻，并约束远期覆盖；近期窗口与分配比例仅在训练／验证阶段确定。完整高频输入（Full-Rate）作为使用更多 Token 的性能与开销参照，不纳入等预算比较，也不预设其精度必然最优。',
76:'在等预算对照基础上，进一步探索运动自适应采样。候选重要性线索包括速度变化、航向变化、加速度及局部曲率，均仅由预测时刻及此前的观测计算。不同量纲的线索需归一化；速度差分与加速度具有相关性，应通过消融确定保留项，避免重复计分。对低速航向不稳定、观测缺失与轨迹噪声设置有效性处理，并在训练／验证数据上确定权重。选点过程结合最低长期覆盖、近期配额与事件分数，去重后严格满足 M 个时间点的预算约束。',
77:'时间模块首先独立验证：固定空间构图规则、解码结构、初始化方式、训练预算和检查点选择准则，仅改变时间选点策略；主对照允许各策略按相同规则训练，另将同权重推理时删点作为鲁棒性诊断。运动特征从同一高频缓冲区按真实时间间隔计算后选取，避免稀疏采样改变差分含义；同一目标及其邻居优先共用选中时间戳，缺失状态使用有效性掩码。随后再评估空间与时间模块的联合效应，判断其收益是否可组合。',
79:'整体模型采用现有并行非自回归多模态解码器作为统一后端，固定输出 20 条候选轨迹。评价时分开报告候选覆盖与单条输出质量：minADE20、minFDE20 分别取候选集合中对应误差的最小值，二者的最优候选可以不同；selected ADE/FDE 对应模型评分最高的同一条轨迹。选择遗憾值（selection regret）定义为该轨迹 ADE 与候选最小 ADE 之差。主实验固定评分机制，使编码与采样收益不与候选选择优化混杂。',
80:'评价同时报告预测精度和实际效率。精度指标明确候选数量、单位、预测时长及漏失率阈值；ETH/UCY 与 SDD 分别沿用其对齐的评估协议，不混合不同坐标单位。效率指标包括参数量、分析器可统计的 FLOPs/MACs、实际边数与槽位利用率、峰值已分配显存、批量大小 1/16/64 的平均及 P95 延迟、吞吐量。时间选择器的打分、选点与数据整理成本计入端到端延迟。ETH/UCY 与 SDD 用于主要精度对比，AV2 行人子集用于高频历史与跨数据集适配验证；适配后结果不表述为零样本泛化能力。',
83:'（2）空间建模目标：以当前基线为基础，完善运动学自适应有向图与时序—社会交叉编码，检验关键邻居保留能力；在同图同权重条件下量化关系计算、显存与延迟变化，确定 Fixed-K 规则布局的收益区间及填充代价。',
94:'（2）Transformer 与稀疏注意力理论：注意力能够表达跨时间与跨智能体依赖，稀疏邻域则限制参与交互的状态集合。本课题将关系编码及交叉注意力的主要计算分别实现为按边 O(ET) 和固定宽度 O(NWT) 路径，同时保留构图、时间自注意力与解码成本的独立分析，避免把局部分支复杂度写成完整模型复杂度。',
99:'稀疏执行先以同图、同权重、同输入验证数值等价，再比较 F16-MaskedDenseExec、F16-EdgeExec 与 F16-FixedK。计时固定设备、软件版本、精度和输入，完成预热与 GPU 同步后交错重复测量，并记录测量波动。时间选择的三种稀疏策略固定历史覆盖、Token 数、训练轮次、初始化与选择准则；Full-Rate 单独作为高输入预算参照。主要结论报告多随机种子均值与标准差，样本差异通过场景级配对重采样评估，避免把同场景多个目标视为完全独立样本。',
100:'AV2 实验采用仅轨迹输入的行人子集诊断协议，与官方多类别预测任务及使用地图的方法在目标、输入和预测设置上存在差异，因此不据此开展直接排名比较。当前高频适配实验用于分析时间输入条件；如开展标准化横比，将另行对齐对应行人基准的数据划分、历史与未来长度、候选数量及地图输入条件。对历史时长同时报告采样率、点数和首末时间戳，避免将 T/f 的名义窗口长度与 (T−1)/f 的采样点间跨度混用。',
108:'三种执行路径保持同一稀疏图语义：F16-MaskedDenseExec 使用完整关系张量并施加掩码；F16-EdgeExec 按选中边构建关系并聚合；F16-FixedK 将邻居整理为 [B,N,W] 的规则槽位，正式配置 W=7，再对 W×T 维度进行连续归一化与归约。无效槽位与填充目标保留相应掩码和边界处理，并报告槽位利用率。图中流程表示信息依赖关系，不代表固定宽度路径完全没有填充计算。最终实现根据数值等价、算术量、显存与稳定计时共同选择。',
112:'如图 2 所示，时间方向从仅包含已观测状态的高频缓冲区取数，在主编码器之前执行选点。第一阶段比较等预算的均匀、近期密集和远期密集采样，同时设置完整 10 Hz 输入参照。除总 Token 数与历史覆盖外，还需统一运动特征的计算方式、时间戳编码和邻居状态对齐。近期密集采样是否有效由多随机种子与分运动类型结果判断，不将已有高频增益直接视为其成立条件。',
113:'第二阶段以经归一化的运动变化线索构造事件分数，在固定预算内分配远期覆盖点、事件点和近期状态。分配变化仅改变预算内各类状态的比例，不随运动复杂度无约束增加输入数量。保留实际时间戳，并在计算速度、加速度时使用真实间隔；对低速转向噪声、重复选点和缺失观测设置确定性处理。自适应策略与静态近期密集策略分别比较，以识别事件选点的额外贡献。',
117:'阶段二：空间稀疏主模型完善（进行中）。已完成 EdgeExec 数值等价、计算量与显存扩展验证，并实现含自环宽度为 7 的 Fixed-K 路径。后续重点复测计时稳定性，补齐数据抽样和运行条件记录，确定最终执行方案及其适用范围。',
118:'阶段三：时间 Token 分配验证。已有 AV2 采样率对照完成相同训练预算及三随机种子复验；下一步开展固定历史覆盖和固定 M 的 Uniform、Recent-Dense、Distant-Dense 对照，并以 Full-Rate 为较高输入预算参照，比较精度、显存与含选点开销的实际延迟。',
124:'（2）技术可行性：项目已具备完整的数据处理、训练、评估及并行解码代码。EdgeExec 与 Fixed-K 均已完成同图同权重的数值等价检查，可直接开展执行效率复测；AV2 采样率对照已完成训练预算控制和多随机种子复验，为固定预算时间选择提供统一实验入口。时间选择的有效性仍需独立验证，不能由现有输入频率实验替代。',
129:'（1）算术量下降但实际延迟不下降。固定宽度布局虽可减少不规则归约，却会引入填充槽位计算；小规模场景还可能受启动和搬运开销主导。拟先在稳定设备状态下交错复测三条路径，并报告中位数、均值及尾部延迟；若规则布局仍未取得稳定收益，则依据实测保留更合适的执行方式，以算术量和显存收益界定轻量化贡献，不将加速作为既成结论。',
140:'（1）基于自适应有向稀疏图与时序—社会交叉注意力的轻量化空间交互方法。利用运动学关系选择有限邻居，并直接建模目标行人与选中邻居历史状态之间的联系；进一步在同图同权重条件下研究按边与固定宽度执行，使邻居选择落实为关系编码和交互分支的计算削减。拟通过图结构消融、数值等价及规模扩展实验，明确预测能力、显存节省与实际延迟的关系及适用边界。',
141:'（2）固定时间 Token 预算下的运动自适应历史采样方法。在保持历史覆盖的前提下，约束进入编码器的时间状态数，先通过等预算静态对照检验近期信息的相对价值，再结合运动事件分数与时间覆盖约束，自适应分配近期状态、事件点和远期趋势点。拟通过真实时间间隔编码及统一训练协议，检验该方法能否以较少历史输入保留高频运动信息，并在计入选择开销后改善精度与效率的权衡。',
}

def replace(p,text):
    props=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear()
    r=p.add_run(text)
    if props is not None: r._r.insert(0,props)

for i,s in changes.items(): replace(ps[i],s)
# Remove duplicate illustrative raster diagrams with stale fixed-neighbor notation.
for p in ps:
    for drawing in list(p._p.xpath('.//w:drawing')):
        ids=drawing.xpath('.//a:blip/@r:embed')
        if ids and Path(str(d.part.related_parts[ids[0]].partname)).name in ['image2.png','image3.png']:
            drawing.getparent().remove(drawing)

abstract='行人轨迹预测需要在多模态运动建模与有限计算资源之间取得平衡。针对空间交互中无关行人对计算较多、历史状态贡献不均衡的问题，本课题拟研究面向时空自适应稀疏建模的轻量化预测方法。空间方向基于相对位置、相对速度及潜在运动冲突构建自适应有向稀疏图，通过时序—社会交叉注意力聚合关键邻居的历史信息，并比较按边与固定宽度执行，检验结构稀疏向实际计算和显存节省的转化条件。时间方向在固定历史覆盖与时间 Token 预算下，比较均匀、近期密集和远期密集采样，进一步结合运动变化与覆盖约束选择关键历史状态。前期空间实验已获得数值等价及算术量、显存收益，实际延迟优势尚待验证；AV2 受控实验支持高频输入的预测价值，但固定预算时间选择仍需独立检验。整体框架采用仅轨迹输入和并行非自回归多模态解码，在 ETH/UCY、SDD 与 AV2 行人子集上分别评估候选覆盖、单条输出误差、计算量、显存及实际延迟，明确两个稀疏模块的独立贡献与联合适用范围。'
replace(d.tables[1].cell(3,1).paragraphs[0],abstract)
rows=[
['稀疏执行等价性','EXP-197，同权重','20 批次；图不一致数 0；轨迹误差≤1.91×10⁻⁶','数值容差内等价','保留边界处理'],
['真实关系计算','ETH/UCY 输入','行人对减少约58%～63%；可统计 FLOPs 减少约9%～10%','算术量下降','不等同实际加速'],
['合成显存扩展','4个非自环邻居+自环','N=64：30.42→4.15 MB；N=128：152.45→9.21 MB','显存扩展改善','限定合成测试条件'],
['实际延迟','EXP-197，B=1/16/64','B=1基本持平；B=64平均延迟增加6.85%','未获稳定加速','稳定环境下复测'],
['固定宽度执行','EXP-198，总宽度7','20批次等价通过；槽位利用率42.7%','实现与等价已验证','计时波动需排查'],
['高频历史对照','EXP-196，3个种子','minADE20：0.2041→0.2001→0.1960 m','支持高频输入价值','固定预算选点待验证'],
['长历史诊断','3.2 s与5.0 s，单种子','minADE20改善1.38%；95%区间跨零','总体收益证据不足','分运动类型复验'],
]
t=d.tables[2]
while len(t.rows)<len(rows)+1: t._tbl.append(deepcopy(t.rows[-1]._tr))
for row,vals in zip(t.rows[1:],rows):
    for cell,s in zip(row.cells,vals): replace(cell.paragraphs[0],s)
replace(d.tables[3].cell(2,1).paragraphs[0],'空间主干与两类稀疏执行已实现；补齐稳定计时及适用范围验证')
replace(d.tables[3].cell(3,1).paragraphs[0],'已有高频多种子对照；开展固定预算的三种采样策略比较')

# Consistent terminology; preserve unchanged references, source layout and identity fields.
for p in d.paragraphs:
    if p.text.startswith('['): continue
    s=p.text.replace('当前已验证的最优主干','当前基线').replace('当前最优 F16 主干','当前 F16 基线').replace('当前最优模型','当前基线模型').replace('现有最优模型','现有基线模型').replace('时间 token','时间 Token').replace('token 预算','Token 预算').replace('。；','；')
    if s!=p.text: replace(p,s)
for p in d.paragraphs:
    if p._p.xpath('.//w:drawing'):
        p.paragraph_format.keep_with_next=True
    if p.text.startswith('图 '):
        p.paragraph_format.keep_with_next=True
    if p.style.name.startswith('Heading'):
        p.paragraph_format.keep_with_next=True
        for r in p.runs: r.font.color.rgb=RGBColor(0,0,0)
for i in [117,118]:
    p=ps[i]
    s=p.text
    replace(p,s)
    for r in p.runs: r.bold=False
for i in [105,106]:
    p=ps[i]
    if not p.text and not p._p.xpath('.//w:drawing'): p._p.getparent().remove(p._p)
# Match the source's label-only emphasis after replacing mixed-format runs.
for p in d.paragraphs:
    s=p.text
    if len(s)>90 and (s.startswith('（') or s.startswith('阶段')):
        boundaries=[s.find(c) for c in ['：','。'] if s.find(c)>3]
        if boundaries:
            cut=min(boundaries)+1
            if cut>=len(s)-1: cut=3
            prop=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
            p.clear()
            for text,bold in [(s[:cut],True),(s[cut:],False)]:
                r=p.add_run(text)
                if prop is not None:r._r.insert(0,deepcopy(prop))
                r.bold=bold
replace(ps[5],'面向时空自适应稀疏建模的\n轻量化行人轨迹预测方法研究')
ps[5].style=d.styles['Title']
ts=d.styles['Title']
ts.font.name='Times New Roman'
ts.font.size=Pt(20)
ts.font.color.rgb=RGBColor(0,0,0)
ts.font.bold=True
ts.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
for el in [ts.element,ps[5]._p]:
    for border in list(el.xpath('.//w:pBdr')): border.getparent().remove(border)
for r in ps[5].runs:
    r.font.color.rgb=RGBColor(0,0,0)
    r.font.name='Times New Roman'
    r._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
ps[43].paragraph_format.first_line_indent=ps[42].paragraph_format.first_line_indent
for t in d.tables:
    for row in t.rows:
        trPr=row._tr.get_or_add_trPr()
        no=OxmlElement('w:cantSplit'); trPr.append(no)
    if t==d.tables[2]:
        hdr=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(hdr)

out=root/'面向时空自适应稀疏建模的轻量化行人轨迹预测方法研究_开题报告_优化版.docx'
d.save(out)
(root/'qa/revision_audit.json').write_text(json.dumps({'changed_paragraphs':list(changes),'sources':['SPARSE_EXECUTION_SERIES/EXP_197_RESULTS','FIXED_K_SPARSE_EXECUTION_SERIES/EXP_198_RESULTS','AV2_SCALING_SERIES/EXP_196B_MULTI_SEED_TEMPORAL/multiseed_summary.json','AV2_SCALING_SERIES/EXP196_RESULT_ANALYSIS.md'],'source_preserved':str(src),'figures':'Retained original simple workflow figures; removed two duplicate illustrative figures containing stale or ambiguous notation.','reference_scope':'Existing reference list preserved; AV2 and Snapshot publication records checked against primary sources. No new references added.'},ensure_ascii=False,indent=2),encoding='utf-8')
print(out)
